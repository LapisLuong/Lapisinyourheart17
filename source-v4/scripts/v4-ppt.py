"""Build a genuinely editable 13-slide PowerPoint; preserve original references."""
from pathlib import Path
import json,math
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR,PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
D=json.loads(Path('reports/v4/ppt-content.json').read_text());prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='FallSense — Giới thiệu ngành Công nghệ thông tin';prs.core_properties.author='Nhóm 06';prs.core_properties.subject='IE005.R11.CNVN — nguyên mẫu hỗ trợ giám sát té ngã'
C={'bg':'EAF7FB','white':'FFFFFF','milk':'FCF9FC','pink':'F2D8E8','lav':'E9DDF8','accent':'C44C99','ink':'20212A','muted':'655D6C','sage':'4F8D61','line':'D8CFE0'}
FONT='Arial'
def rgb(c):return RGBColor.from_string(C.get(c,c))
def box(s,x,y,w,h,fill='white',radius=True,line=None):
 q=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=rgb(fill)
 if line:q.line.color.rgb=rgb(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if radius:
  try:q.adjustments[0]=.12
  except Exception:pass
 q._element.spPr.append(OxmlElement('a:effectLst'))
 try:q._element.xpath('./p:style/a:effectRef')[0].set('idx','0')
 except Exception:pass
 return q
def txt(s,text,x,y,w,h,size=20,bold=False,color='ink',align=None):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));t=q.text_frame;t.clear();t.word_wrap=True;t.margin_left=t.margin_right=0;t.margin_top=Pt(2);t.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=t.paragraphs[0] if i==0 else t.add_paragraph();p.text=line;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(6)
  if align is not None:p.alignment=align
 return q
def line(s,a,b,color='accent',width=2):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(a[0]),Inches(a[1]),Inches(b[0]),Inches(b[1]));q.line.color.rgb=rgb(color);q.line.width=Pt(width);return q
def circle(s,x,y,d,fill='pink'):
 q=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(d),Inches(d));q.fill.solid();q.fill.fore_color.rgb=rgb(fill);q.line.fill.background();return q
def chip(s,text,x,y,w=2.0,fill='pink'):
 box(s,x,y,w,.38,fill);txt(s,text,x+.13,y+.03,w-.26,.32,11,True,'accent')
def mascot(s,x,y,w=2.8):
 pts=[(0,120),(200,75),(500,0),(800,75),(1000,120),(970,550),(800,850),(500,1080),(200,850),(30,550)]
 b=s.shapes.build_freeform(0,120,scale=Inches(w)/1000);b.add_line_segments(pts[1:],close=True);q=b.convert_to_shape(Inches(x),Inches(y));q.fill.gradient();q.fill.gradient_stops[0].color.rgb=rgb('F9E6F2');q.fill.gradient_stops[1].color.rgb=rgb('DDA0C6');q.line.color.rgb=rgb('FFFFFF');q.line.width=Pt(3)
 for a in [.32,.68]:
  circle(s,x+w*a-w*.035,y+w*.40,w*.07,'ink');circle(s,x+w*a-w*.018,y+w*.405,w*.018,'white')
  q=circle(s,x+w*a-w*.065,y+w*.55,w*.13,'EAA8C9');q.height=Inches(w*.05)
 for a in [-.06,.94]:
  q=circle(s,x+w*a,y+w*.50,w*.14,'pink');q.height=Inches(w*.31)
 line(s,(x+w*.44,y+w*.60),(x+w*.50,y+w*.64),'96617F',2.2);line(s,(x+w*.50,y+w*.64),(x+w*.56,y+w*.60),'96617F',2.2)
 line(s,(x+w*.42,y+w*.79),(x+w*.49,y+w*.85),'white',3);line(s,(x+w*.49,y+w*.85),(x+w*.62,y+w*.71),'white',3)
 for a in [.28,.62]:circle(s,x+w*a,y+w*1.02,w*.13,'pink')
def new(i):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb('bg')
 # A restrained native background panel, not a raster slide screenshot.
 box(s,10.7,0,2.64,7.5,'lav',False);box(s,0,7.18,13.34,.32,'pink',False)
 txt(s,'FALLSENSE  /  A LITTLE GUARDIAN',.60,.32,9,.30,10,True,'accent');txt(s,f'{i+1:02d} / 13',11.55,.32,1.1,.3,10,True,'muted',PP_ALIGN.RIGHT)
 title=D['slides'][i]['title'].replace('\n',' ')
 if i>0:txt(s,title,.60,.95,11.9,1.0,32,True)
 txt(s,D['slides'][i]['source'],.60,6.83,11.9,.30,10,False,'muted')
 notes=s.notes_slide.notes_text_frame;notes.text=D['slides'][i]['note']+'\n\nNguồn: '+D['slides'][i]['source']+'\n\nBản đồ nguồn: slide 1–8 = bản gốc 1–8 đã sửa; slide 9 mới; slide 10 = gốc 9; slide 11 mới; slide 12 = gốc 10; slide 13 mới.'
 if i==1:notes.text+='\nhttps://www.who.int/news-room/fact-sheets/detail/falls\nhttps://doi.org/10.1136/bmj.a2227'
 if i in [4,11]:notes.text+='\n\nXung đột cần đối chiếu: ảnh Live Monitor ghi vạch 0,60, tài liệu góp ý nêu 0,68. Không xem ảnh chụp là bằng chứng ngưỡng trong mã đang triển khai. Chưa có app source để giải quyết.'
 if i==6:notes.text+='\nhttps://ai.google.dev/edge/mediapipe/solutions/vision/pose_landmarker'
 # Native click-advance fade, no forced auto-advance. This is not video playback.
 tr=OxmlElement('p:transition');tr.set('spd','med');tr.append(OxmlElement('p:fade'));s._element.append(tr)
 return s

def cards(s,items,y=2.45):
 width=(12.13-.25*(len(items)-1))/len(items)
 for i,t in enumerate(items):
  x=.60+i*(width+.25);parts=t.split('\n');box(s,x,y,width,2.65,'white');txt(s,parts[0],x+.23,y+.26,width-.46,.95,23,True,'accent' if i==0 else 'ink');txt(s,'\n'.join(parts[1:]),x+.23,y+1.30,width-.46,1.13,18,False,'muted')

for i,item in enumerate(D['slides']):
 s=new(i)
 if i==0:
  chip(s,'GIỚI THIỆU NGÀNH CÔNG NGHỆ THÔNG TIN',.6,1.13,4.9)
  txt(s,'FallSense',.60,1.85,8,1.08,62,True)
  txt(s,'Hệ thống hỗ trợ giám sát và\ncảnh báo té ngã bằng camera',.66,3.02,8.5,1.15,29,True,'accent')
  txt(s,'Nhóm 06 – IE005.R11.CNVN\nGiảng viên: TS. Nguyễn Tấn Hoàng Phước\nTrường Đại học Công nghệ Thông tin – ĐHQG-HCM',.66,4.67,9.1,1.40,18,False,'muted')
  mascot(s,9.52,2.15,2.65)
 elif i==1:
  for j,(value,label) in enumerate([('684.000','ca tử vong do té ngã/năm\nƯớc tính toàn cầu · WHO 2021'),('53/66','không tự đứng dậy\nSau ít nhất một lần ngã'),('20/66','từng nằm sàn ≥ 1 giờ\nTrong nhóm đã ngã')]):
   x=.60+j*4.12;box(s,x,2.4,3.87,2.65,'white');txt(s,value,x+.25,2.66,3.4,.8,44,True,'accent');txt(s,label,x+.25,3.68,3.37,1.1,19,False,'muted')
  txt(s,'53/66 và 20/66 thuộc nghiên cứu người trên 90 tuổi đã ngã; không phải tỷ lệ dân số chung.',.70,5.55,11.85,.68,18,True)
 elif i==2:
  cards(s,item['items']);chip(s,'NGUYÊN MẪU HỌC TẬP',.65,5.62,3.2);txt(s,'Không chẩn đoán y khoa · Không tự gọi cấp cứu',4.1,5.65,8,.5,19,True)
 elif i==3:
  for j,(heading,sub) in enumerate([('Camera','Hình ảnh'),('MediaPipe','Điểm mốc tư thế'),('Fall Score','Đặc trưng chuyển động'),('Trạng thái','Diễn biến thời gian'),('Phản hồi','Giao diện / cảnh báo')]):
   x=.6+j*2.48;box(s,x,2.7,2.18,2.0,'white');circle(s,x+.22,2.91,.36,'pink');txt(s,str(j+1),x+.31,2.94,.20,.25,12,True,'accent');txt(s,heading,x+.2,3.45,1.8,.45,20,True);txt(s,sub,x+.2,4.02,1.8,.5,15,False,'muted')
   if j<4:txt(s,'→',x+2.2,3.35,.30,.6,25,True,'accent')
  txt(s,'Một khung hình chưa đủ — cần chất lượng đầu vào và diễn biến theo thời gian.',.65,5.55,12,.7,20,True)
 elif i==4:
  s.shapes.add_picture('public/assets/live-monitor.jpg',Inches(.6),Inches(2.0),width=Inches(8.65))
  labels=['Camera / mô phỏng','Khung xương','ID theo dõi tạm thời','Fall Score: điểm 0–1','Trạng thái hiện tại']
  for j,t in enumerate(labels):
   y=2.1+j*.77;circle(s,9.55,y,.38,'accent');txt(s,str(j+1),9.66,y+.03,.2,.3,13,True,'white');txt(s,t,10.03,y,2.65,.65,17,True)
  for j,(xx,yy) in enumerate([(.30,.25),(.49,.48),(.79,.25),(.81,.38),(.89,.25)]):
   x=.6+8.65*xx;y=2.0+(8.65*1041/2047)*yy;circle(s,x,y,.25,'accent');txt(s,str(j+1),x+.07,y+.005,.15,.25,10,True,'white')
  chip(s,'DEMO SIM',9.58,6.10,2.45);txt(s,'Không phải phép đo hiệu quả phát hiện.',.66,6.47,8.7,.3,13,False,'muted')
 elif i==5:
  cards(s,['FALLEN\nMở hộp thoại theo góp ý','15 giây\n“Tôi ổn” để hủy hoặc chọn gửi ngay','Thử gửi email\nCần dịch vụ và kết nối mạng'])
  txt(s,'Không mô tả hai chu kỳ 15 giây nối tiếp; chưa kiểm chứng phiên ứng dụng hiện hành.',.66,5.72,11.95,.72,18,True,'accent')
 elif i==6:
  coords=[(2.25,2.65),(4.0,2.65),(2.45,4.05),(3.8,4.05),(2.25,5.1),(4.,5.1),(2.15,6.12),(4.1,6.12)]
  for a,b in [(0,1),(0,2),(1,3),(2,3),(2,4),(3,5),(4,6),(5,7)]:line(s,coords[a],coords[b],'sage',3)
  circle(s,2.75,1.99,.72,'pink')
  for j,((x,y),num) in enumerate(zip(coords,[11,12,23,24,25,26,27,28])):circle(s,x-.06,y-.06,.12,'accent');txt(s,str(num),x+.14,y-.15,.45,.4,16,True,'accent')
  for j,t in enumerate(['33 điểm mốc là đầu vào tư thế','Vai 11/12 · Hông 23/24\nGối 25/26 · Cổ chân 27/28','Visibility khác Fall Score\nMediaPipe không tự kết luận ngã']):
   box(s,6.1,2.25+j*1.25,6.6,1.02,'white');txt(s,t,6.35,2.43+j*1.25,6.05,.75,20,True if j==0 else False)
 elif i==7:
  cards(s,item['items']);txt(s,'Chỉ đánh giá khi dữ liệu thân đủ tin cậy; không đánh đồng “thiếu dữ liệu” với “an toàn”.',.66,5.60,11.8,.8,19,True)
 elif i==8:
  for j,(r,label,color) in enumerate([(.55,'Dọc: r < 0,75','sage'),(1.,'Gần vuông','muted'),(1.6,'Ngang: r > 1,30','accent')]):
   x=.6+j*4.12;box(s,x,2.24,3.87,2.62,'white');h=1.3 if j==0 else 1.0 if j==1 else .9;w=h*r;xx=x+(3.87-w)/2;yy=2.5+(1.4-h)/2;box(s,xx,yy,w,h,'pink',False,'accent');txt(s,label,x+.2,4.12,3.45,.5,19,True,color,PP_ALIGN.CENTER)
  txt(s,'W = xmax − xmin    H = ymax − ymin    r = W/H',.67,5.15,11.8,.6,25,True)
  txt(s,'Tọa độ chuẩn hóa; khung ngang chỉ là một dấu hiệu. Ngưỡng dẫn theo góp ý, chưa đối chiếu source.',.67,5.91,11.8,.65,17,False,'muted')
 elif i==9:
  for j,(name,w) in enumerate(zip(D['featureNames'],D['weights'])):
   y=2.1+j*.55;txt(s,name,.65,y,3.6,.42,17,False);box(s,4.28,y+.10,4.3,.18,'white',False);box(s,4.28,y+.10,4.3*w/.22,.18,'accent' if j%2 else 'B09ACB',False);txt(s,f'{w:.2f}',8.85,y,.9,.4,17,True)
  box(s,10.0,2.13,2.7,3.85,'white');txt(s,'Σ wi =',10.25,2.5,2.2,.5,26,True);txt(s,'1,00',10.25,3.16,2.2,.9,44,True,'accent');txt(s,'Điểm 0–1\nKhông phải xác suất',10.25,4.42,2.15,1.15,19,True)
  txt(s,'FallScore = clamp(mROI × Σ wi si)',.67,6.15,11.9,.5,22,True)
 elif i==10:
  chip(s,'VÍ DỤ GIẢ ĐỊNH',.63,1.91,2.85)
  table=s.shapes.add_table(8,4,Inches(.63),Inches(2.5),Inches(9.2),Inches(3.60)).table
  for col,w in zip(table.columns,[3.7,1.8,1.5,2.2]):col.width=Inches(w)
  rows=[['Thành phần','Điểm si','Trọng số','Đóng góp']]+[[n,f'{a:.2f}',f'{b:.2f}',f'{a*b:.3f}'] for n,a,b in zip(D['featureNames'],D['example'],D['weights'])]
  for r,row in enumerate(rows):
   for c,text in enumerate(row):
    cell=table.cell(r,c);cell.text=text;cell.fill.solid();cell.fill.fore_color.rgb=rgb('pink' if r==0 else 'white');cell.margin_left=Inches(.10);cell.margin_top=Inches(.05);cell.margin_bottom=0
    for p in cell.text_frame.paragraphs:p.font.name=FONT;p.font.size=Pt(15 if c==0 else 16);p.font.bold=r==0;p.font.color.rgb=rgb('ink')
  txt(s,'0,521',10.07,3.05,2.75,1.0,42,True,'accent');txt(s,'Giả sử mROI = 1\nChưa dùng model hỗ trợ',10.07,4.22,2.6,1.2,17,False,'muted');txt(s,'Không quy đổi thành độ chính xác 52,1%. Còn xét điều kiện hình học và thời gian.',.65,6.27,11.8,.4,16,True)
 elif i==11:
  for j,state in enumerate(['STANDING','PRE_FALL','FALLEN','ALERT']):
   x=.65+j*3.12;box(s,x,2.50,2.73,.88,'pink' if j==2 else 'white');txt(s,state,x+.1,2.76,2.53,.42,20,True,'accent' if j==2 else 'ink',PP_ALIGN.CENTER)
   if j<3:txt(s,'→',x+2.78,2.64,.35,.55,27,True,'accent')
  recovery=s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW,Inches(.75),Inches(3.48),Inches(11.55),Inches(.13));recovery.fill.solid();recovery.fill.fore_color.rgb=rgb('sage');recovery.line.fill.background()
  chip(s,'Phục hồi / Tôi ổn — theo điều kiện',4.45,3.36,4.5,'milk')
  cond=[('Nghi ngờ','Hạ nhanh + thân nghiêng,\nhoặc điểm ≥ 0,45\ntích lũy đủ bước.'),('Xác nhận','Dấu hiệu ngang tích lũy,\nhoặc điểm ≥ 0,68 + góc >60°,\nhoặc giảm cao nhanh + gần đáy.'),('Leo thang','Bất động tích lũy\nđạt ngưỡng cấu hình,\nmặc định 15 giây.')]
  for j,(a,b) in enumerate(cond):
   x=.65+j*4.15;box(s,x,3.85,3.9,1.72,'white');txt(s,a,x+.20,4.0,3.5,.4,19,True,'accent');txt(s,b,x+.20,4.48,3.5,.95,15,False,'muted')
  txt(s,'NO_PERSON / PARTIAL_POSE: chất lượng đầu vào · Bộ đếm bước khác thời gian giây.',.66,1.94,11.85,.42,17,True)
  txt(s,'Theo góp ý: hộp thoại mở từ FALLEN. Điều kiện rút gọn; chưa kiểm tra source hiện hành.',.66,6.10,11.9,.54,16,True,'accent')
 elif i==12:
  cards(s,['Đề xuất kiểm thử\nĐứng, cúi, thiếu điểm mốc,\nngã mô phỏng có kiểm soát.','Ghi bằng chứng\nPhát hiện đúng / bỏ sót,\nbáo sai, modal và email.','Nhận diện giới hạn\nÁnh sáng, che khuất, góc máy;\nquyền camera và mạng.'])
  txt(s,'Chưa có số đo hiệu quả thực tế trong bộ tài liệu. Test phần mềm không chứng minh độ chính xác lâm sàng.',.67,5.62,11.85,.8,18,True,'accent')

assert len(prs.slides)==13
out=Path('out/FallSense_Academic_13_Slides_Editable.pptx');prs.save(out)
# Structural / editability QA; content is native text, shapes and tables, not slide screenshots.
from zipfile import ZipFile
with ZipFile(out) as z:
 assert z.testzip() is None
 assert sum(1 for n in z.namelist() if __import__('re').fullmatch(r'ppt/slides/slide\d+\.xml',n))==13
 assert sum(1 for n in z.namelist() if __import__('re').fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',n))==13
assert abs(sum(D['weights'])-1)<1e-9
assert abs(sum(a*b for a,b in zip(D['weights'],D['example']))-.521)<1e-9
print('Created',out,'13 editable slides + notes; arithmetic and ZIP structure passed')
